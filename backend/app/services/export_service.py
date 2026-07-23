import os
import time
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from app.config import settings
from app.models.paper import ResearchPaper

class ExportService:
    def export_to_excel(self, papers: List[ResearchPaper]) -> str:
        """Generates a styled Excel (.xlsx) spreadsheet report of research publications."""
        filename = f"promptflow_ai_research_report_{int(time.time())}.xlsx"
        filepath = os.path.join(settings.EXPORTS_DIR, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Research Publications"

        # Title block
        ws.merge_cells("A1:H1")
        ws["A1"] = "promptFlow-AI - Publications Audit Report"
        ws["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 40

        # Subtitle
        ws.merge_cells("A2:H2")
        ws["A2"] = f"Generated on {time.strftime('%Y-%m-%d %H:%M:%S')} | Total Papers: {len(papers)}"
        ws["A2"].font = Font(name="Arial", size=11, italic=True, color="475569")
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 24

        # Table Headers
        headers = ["ID", "Title", "Authors", "Department", "Journal / Conference", "Year", "DOI / Link", "ML Dept Confidence"]
        ws.append([]) # blank row 3
        ws.append(headers) # row 4
        ws.row_dimensions[4].height = 28

        header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_thin

        # Data rows
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        for idx, paper in enumerate(papers, start=5):
            row_data = [
                paper.id,
                paper.title,
                paper.authors_raw or ", ".join(paper.authors_list or []),
                paper.department or "General",
                paper.journal_or_conference or "N/A",
                paper.publication_year or 2026,
                paper.doi or "N/A",
                f"{round((paper.ml_confidence or 0.0) * 100, 1)}%"
            ]
            ws.append(row_data)
            ws.row_dimensions[idx].height = 22

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=idx, column=col_num)
                cell.font = Font(name="Arial", size=10)
                cell.border = border_thin
                if idx % 2 == 0:
                    cell.fill = alt_fill
                if col_num in [1, 6, 8]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row > 2 and cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        wb.save(filepath)
        return filepath

    def export_to_pdf(self, papers: List[ResearchPaper]) -> str:
        """Generates a publication-ready PDF summary report of research papers."""
        filename = f"promptflow_ai_research_report_{int(time.time())}.pdf"
        filepath = os.path.join(settings.EXPORTS_DIR, filename)

        doc = SimpleDocTemplate(filepath, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=20,
            textColor=colors.HexColor('#1E3A8A'),
            spaceAfter=6,
            alignment=1 # Center
        )
        sub_style = ParagraphStyle(
            'ReportSub',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=10,
            textColor=colors.HexColor('#64748B'),
            spaceAfter=15,
            alignment=1
        )
        body_style = ParagraphStyle(
            'TableBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            textColor=colors.HexColor('#1E293B'),
            leading=11
        )
        header_style = ParagraphStyle(
            'TableHeader',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=9,
            textColor=colors.white,
            alignment=1
        )

        elements = []
        elements.append(Paragraph("promptFlow-AI System", title_style))
        elements.append(Paragraph(f"Academic Research Publications Audit | Generated: {time.strftime('%Y-%m-%d')} | Total Records: {len(papers)}", sub_style))
        elements.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor('#3B82F6'), spaceAfter=15))

        # Table Setup
        data = [[
            Paragraph("Title", header_style),
            Paragraph("Authors", header_style),
            Paragraph("Dept", header_style),
            Paragraph("Journal / Conf", header_style),
            Paragraph("Year", header_style)
        ]]

        for p in papers:
            title_p = Paragraph(p.title[:80] + ("..." if len(p.title) > 80 else ""), body_style)
            auth_str = p.authors_raw or ", ".join(p.authors_list or [])
            auth_p = Paragraph(auth_str[:60] + ("..." if len(auth_str) > 60 else ""), body_style)
            dept_p = Paragraph(p.department or "General", body_style)
            journ_p = Paragraph((p.journal_or_conference or "N/A")[:50], body_style)
            year_p = Paragraph(str(p.publication_year or 2026), body_style)
            data.append([title_p, auth_p, dept_p, journ_p, year_p])

        col_widths = [160, 130, 90, 110, 40]
        t = Table(data, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E3A8A')),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')])
        ]))

        elements.append(t)
        doc.build(elements)
        return filepath

export_service = ExportService()
